"""Unit tests using real authored Office/PDF bytes and the production parsers.

Install the ``office`` extra to run every format. No parser or reader is mocked.
All generated documents, snapshots and evidence receipts stay under tmp_path.
"""

import hashlib
import json
import tomllib
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

import pytest

from reward_engine.adapters.workbuddy.evidence import collect, inspect_file
from reward_engine.adapters.workbuddy.measurement import (
    WorkBuddyTools,
    workspace_records,
)
from reward_engine.core.outcome.context import build_context
from reward_engine.core.outcome.harness import has_product_evidence

FORMATS = ("xlsx", "docx", "pptx", "pdf")
AUTHORED = "Authored delivery evidence"
SOURCE_LABEL = "Unmodified public workspace files"
PRODUCT_LABEL = "Created or modified workspace files (actual delivery evidence)"


def make_xlsx(path):
    openpyxl = pytest.importorskip("openpyxl", reason="requires the office extra")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Budget α"
    for row in [
        ["Item", "Amount"],
        [AUTHORED, 1200],
        ["Services", 300],
        ["Total", "=SUM(B2:B3)"],
        ["Reviewed", False],
        ["Paid", 0],
        ["Due", datetime(2026, 10, 1)],
    ]:
        sheet.append(row)
    sheet["A9"] = "Merged authored note"
    sheet.merge_cells("A9:B9")
    source = book.create_sheet("Source")
    source["A1"] = "Original source assumption"
    source.sheet_state = "hidden"
    book.save(path)
    book.close()


def make_docx(path):
    docx = pytest.importorskip("docx", reason="requires the office extra")
    doc = docx.Document()
    doc.add_heading("Delivery report", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("Authored ").bold = True
    paragraph.add_run("delivery evidence")
    doc.add_paragraph("Unicode evidence: 证据 α\nSecond authored line")
    table = doc.add_table(rows=3, cols=2)
    for cells, values in zip(
        table.rows,
        [("Metric", "Value"), ("Revenue", "1500"), ("Status", "Reviewed\nAccepted")],
        strict=True,
    ):
        for cell, value in zip(cells.cells, values, strict=True):
            cell.text = value
    doc.save(path)


def make_pptx(path):
    pptx = pytest.importorskip("pptx", reason="requires the office extra")
    from pptx.util import Inches

    deck = pptx.Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    title = first.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    title.name = "Authored title"
    title.text = AUTHORED
    body = first.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(2))
    body.name = "Authored body"
    body.text = "Revenue: 1500\nEvidence: 证据 α"
    first.notes_slide.notes_text_frame.text = "Speaker note: check the source table"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    conclusion = second.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    conclusion.name = "Conclusion"
    conclusion.text = "Second slide conclusion"
    deck.save(path)


def make_pdf(path):
    pypdf = pytest.importorskip("pypdf", reason="requires the office extra")
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    # Real PDF text streams with a standard font; no external renderer required.
    writer = pypdf.PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for lines in [
        (AUTHORED, "First page revenue: 1500"),
        ("Second page conclusion", "Source checked"),
    ]:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(
            (
                "BT /F1 12 Tf 72 720 Td 18 TL "
                + " T* ".join(f"({line}) Tj" for line in lines)
                + " ET"
            ).encode("ascii")
        )
        page[NameObject("/Contents")] = stream
    writer.write(path)
    writer.close()


BUILDERS = {"xlsx": make_xlsx, "docx": make_docx, "pptx": make_pptx, "pdf": make_pdf}


@pytest.fixture(params=FORMATS)
def office_file(request, tmp_path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    path = workspace / ("delivery." + request.param)
    BUILDERS[request.param](path)
    return path


def evidence_for(candidate):
    _, evidence = build_context(
        {"task_id": "office-validation", "prompt": "Read the actual Office delivery"},
        [candidate],
    )
    return evidence


def test_xlsx_preserves_coordinates_formulas_types_and_hidden_sheets(tmp_path):
    path = tmp_path / "delivery.xlsx"
    make_xlsx(path)
    result = inspect_file(path)
    assert [sheet["name"] for sheet in result["sheets"]] == ["Budget α", "Source"]
    cells = {cell["cell"]: cell for cell in result["sheets"][0]["cells"]}
    assert cells["A2"] == {
        "cell": "A2",
        "value_or_formula": AUTHORED,
        "cached_value": AUTHORED,
    }
    assert cells["B2"]["value_or_formula"] == cells["B2"]["cached_value"] == 1200
    assert cells["B4"] == {
        "cell": "B4",
        "value_or_formula": "=SUM(B2:B3)",
        "cached_value": None,
    }
    assert cells["B5"]["value_or_formula"] is cells["B5"]["cached_value"] is False
    assert (
        type(cells["B6"]["value_or_formula"]) is int
        and cells["B6"]["value_or_formula"] == 0
    )
    assert (
        cells["B7"]["value_or_formula"]
        == cells["B7"]["cached_value"]
        == datetime(2026, 10, 1)
    )
    assert cells["A9"]["value_or_formula"] == "Merged authored note"
    assert "B9" not in cells and "A8" not in cells
    assert result["sheets"][1]["cells"] == [
        {
            "cell": "A1",
            "value_or_formula": "Original source assumption",
            "cached_value": "Original source assumption",
        }
    ]


def test_xlsx_reads_stored_formula_cache_without_recalculating(tmp_path):
    path = tmp_path / "cached.xlsx"
    make_xlsx(path)
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    # Author a deliberately stale cached value in the actual XLSX XML, just as
    # a spreadsheet producer can. The reader must preserve it, not recompute it.
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    sheet = ElementTree.fromstring(entries["xl/worksheets/sheet1.xml"])
    cell = sheet.find(".//s:c[@r='B4']", namespace)
    assert cell is not None
    cell.find("s:v", namespace).text = "1499"
    entries["xl/worksheets/sheet1.xml"] = ElementTree.tostring(sheet)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    result = inspect_file(path)
    formula = next(
        cell for cell in result["sheets"][0]["cells"] if cell["cell"] == "B4"
    )
    assert formula == {
        "cell": "B4",
        "value_or_formula": "=SUM(B2:B3)",
        "cached_value": 1499,
    }


def test_docx_extracts_authored_paragraphs_and_table_cells(tmp_path):
    path = tmp_path / "delivery.docx"
    make_docx(path)
    assert inspect_file(path) == {
        "paragraphs": [
            "Delivery report",
            AUTHORED,
            "Unicode evidence: 证据 α\nSecond authored line",
        ],
        "tables": [
            [["Metric", "Value"], ["Revenue", "1500"], ["Status", "Reviewed\nAccepted"]]
        ],
    }


def test_pptx_extracts_slide_order_text_geometry_and_speaker_notes(tmp_path):
    path = tmp_path / "delivery.pptx"
    make_pptx(path)
    result = inspect_file(path)
    assert result["slide_count"] == 2
    assert [slide["slide"] for slide in result["slides"]] == [1, 2]
    first, second = result["slides"]
    assert first["shapes"] == [
        {
            "name": "Authored title",
            "text": AUTHORED,
            "left": 914400,
            "top": 914400,
            "width": 6400800,
            "height": 914400,
        },
        {
            "name": "Authored body",
            "text": "Revenue: 1500\nEvidence: 证据 α",
            "left": 914400,
            "top": 1828800,
            "width": 6400800,
            "height": 1828800,
        },
    ]
    assert first["notes"] == "Speaker note: check the source table"
    assert second["shapes"][0]["text"] == "Second slide conclusion"
    assert second["notes"] == ""
    assert result["visual_quality"] == "unmeasured"


def test_pdf_extracts_authored_text_in_page_order(tmp_path):
    path = tmp_path / "delivery.pdf"
    make_pdf(path)
    result = inspect_file(path)
    assert [page["page"] for page in result["pages"]] == [1, 2]
    assert result["pages"][0]["text"].splitlines() == [
        AUTHORED,
        "First page revenue: 1500",
    ]
    assert result["pages"][1]["text"].splitlines() == [
        "Second page conclusion",
        "Source checked",
    ]
    assert result["visual_quality"] == "unmeasured"


def test_format_dispatch_is_case_insensitive_and_does_not_change_bytes(office_file):
    before = office_file.read_bytes()
    expected = inspect_file(office_file)
    renamed = office_file.with_suffix(office_file.suffix.upper())
    office_file.rename(renamed)
    assert inspect_file(renamed) == expected
    assert renamed.read_bytes() == before


@pytest.mark.asyncio
async def test_real_office_delivery_provenance_survives_collection_and_read(
    office_file, tmp_path
):
    before = office_file.read_bytes()
    digest = hashlib.sha256(before).hexdigest()
    output = tmp_path / "receipts"
    candidate, hashes = collect(
        office_file.parent,
        candidate_id="actual-office",
        final_response="A claim, not a parser result",
        initial_hashes={},
        output=output,
    )
    assert hashes == {office_file.name: digest}
    assert candidate["environment"] == {
        "workspace": "/workspace",
        "file_count": 1,
        "deleted_files": [],
        "snapshot_phase": "before_official_verifier",
        "visual_quality": "unmeasured",
    }
    assert candidate["artifact"]["changed_files"] == [
        {
            "path": office_file.name,
            "sha256": digest,
            "bytes": len(before),
            "changed": True,
        }
    ]
    assert json.loads((output / "candidate.json").read_text()) == candidate
    assert (
        json.loads((output / "inventory.json").read_text())
        == candidate["artifact"]["changed_files"]
    )
    evidence = evidence_for(candidate)
    (record,) = workspace_records(evidence)
    assert record["sha256"] == digest and record["bytes"] == len(before)
    assert record["changed"] is True
    # Datetimes in XLSX are intentionally serialized by the evidence boundary.
    assert record["content"] == json.loads(
        json.dumps(inspect_file(office_file), default=str)
    )
    assert AUTHORED in json.dumps(record["content"], ensure_ascii=False)
    for ref in record["evidence_refs"]:
        assert evidence[ref]["source"] == PRODUCT_LABEL
        assert evidence[ref]["excerpt_only"] is False
    receipt = await WorkBuddyTools().run(
        {"tool": "workspace.read", "arguments": {"paths": [office_file.name]}},
        evidence,
    )
    assert receipt["status"] == "observed"
    body = json.loads(receipt["result"]["text"])
    assert body == {"files": [record], "missing_paths": []}
    assert receipt["result"]["observed_product_refs"] == record["evidence_refs"]
    assert receipt["result"]["next_offset"] is None
    assert has_product_evidence(["tool:read"], evidence, {"tool:read": receipt})
    assert office_file.read_bytes() == before


@pytest.mark.asyncio
async def test_unchanged_office_source_is_not_delivery_read_evidence(office_file):
    _, initial = collect(office_file.parent, candidate_id="initial", final_response="")
    candidate, hashes = collect(
        office_file.parent,
        candidate_id="unchanged",
        final_response="I authored everything",
        initial_hashes=initial,
    )
    assert hashes == initial and candidate["artifact"]["changed_files"] == []
    evidence = evidence_for(candidate)
    (record,) = workspace_records(evidence)
    assert record["changed"] is False
    assert all(
        evidence[ref]["source"] == SOURCE_LABEL for ref in record["evidence_refs"]
    )
    receipt = await WorkBuddyTools().run(
        {"tool": "workspace.read", "arguments": {"paths": [office_file.name]}},
        evidence,
    )
    assert receipt["status"] == "observed"
    assert AUTHORED in receipt["result"]["text"]
    assert "observed_product_refs" not in receipt["result"]
    assert not has_product_evidence(["tool:read"], evidence, {"tool:read": receipt})


@pytest.mark.asyncio
async def test_corrupt_office_bytes_retain_hash_and_do_not_claim_product_observation(
    office_file,
):
    original = office_file.read_bytes()
    office_file.write_bytes(original[:16])
    candidate, hashes = collect(
        office_file.parent,
        candidate_id="broken",
        final_response="Complete",
        initial_hashes={},
    )
    assert hashes == {office_file.name: hashlib.sha256(original[:16]).hexdigest()}
    evidence = evidence_for(candidate)
    (record,) = workspace_records(evidence)
    assert record["content"]["status"] == "unknown"
    assert record["content"]["error_type"]
    assert record["bytes"] == 16 and record["changed"] is True
    for tool in ("workspace.read", "workspace.inspect"):
        receipt = await WorkBuddyTools().run(
            {"tool": tool, "arguments": {"paths": [office_file.name]}}, evidence
        )
        assert receipt["status"] == "observed"
        assert "observed_product_refs" not in receipt["result"]
        assert not has_product_evidence(
            ["tool:broken"], evidence, {"tool:broken": receipt}
        )
        body = json.loads(receipt["result"]["text"])
        content_key = "content" if tool == "workspace.read" else "parse_error"
        assert body["files"][0][content_key] == record["content"]


def test_office_and_runtime_extras_supply_every_supported_parser():
    config = tomllib.loads(
        (Path(__file__).resolve().parents[2] / "pyproject.toml").read_text()
    )
    expected = {"openpyxl", "python-docx", "python-pptx", "pypdf"}
    for extra in ("office", "runtime"):
        requirements = config["project"]["optional-dependencies"][extra]
        names = {requirement.split(">=")[0] for requirement in requirements}
        assert expected <= names, f"{extra} does not supply every Office/PDF parser"
    assert config["project"]["requires-python"] == ">=3.11"


@pytest.mark.parametrize("kind", ["table", "group"])
@pytest.mark.asyncio
async def test_pptx_preserves_authored_table_and_group_text(tmp_path, kind):
    pptx = pytest.importorskip("pptx", reason="requires the office extra")
    from pptx.util import Inches

    path = tmp_path / "structured.pptx"
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    marker = f"Authored {kind} evidence: total 1500"
    if kind == "table":
        shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2))
        shape.table.cell(0, 0).text = "Metric"
        shape.table.cell(1, 0).text = marker
    else:
        group = slide.shapes.add_group_shape()
        shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        shape.text = marker
    deck.save(path)
    # Reopen the actual deck to prove that the authoring library persisted text.
    actual = pptx.Presentation(path).slides[0].shapes[0]
    persisted = (
        actual.table.cell(1, 0).text if kind == "table" else actual.shapes[0].text
    )
    assert persisted == marker
    result = inspect_file(path)
    extracted = result["slides"][0]["shapes"][0]
    assert extracted["text"] == ""
    if kind == "table":
        assert extracted["table"] == [["Metric", ""], [marker, ""]]
    else:
        assert extracted["shapes"][0]["text"] == marker
    candidate, hashes = collect(
        tmp_path, candidate_id="structured", final_response="", initial_hashes={}
    )
    evidence = evidence_for(candidate)
    (record,) = workspace_records(evidence)
    assert record["content"] == result
    assert (
        record["sha256"]
        == hashes[path.name]
        == hashlib.sha256(path.read_bytes()).hexdigest()
    )
    receipt = await WorkBuddyTools().run(
        {"tool": "workspace.read", "arguments": {"paths": [path.name]}}, evidence
    )
    assert receipt["status"] == "observed"
    assert json.loads(receipt["result"]["text"])["files"][0]["content"] == result
    assert receipt["result"]["observed_product_refs"] == record["evidence_refs"]
    assert has_product_evidence(
        ["tool:structured"], evidence, {"tool:structured": receipt}
    )


def test_office_snapshot_tracks_modified_created_unchanged_and_deleted_files(tmp_path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    report = workspace / "report.docx"
    reference = workspace / "reference.pdf"
    removed = workspace / "obsolete.xlsx"
    make_docx(report)
    make_pdf(reference)
    make_xlsx(removed)
    original, initial = collect(workspace, candidate_id="original", final_response="")
    docx = pytest.importorskip("docx", reason="requires the office extra")
    doc = docx.Document(report)
    doc.add_paragraph("Revised authored conclusion")
    doc.save(report)
    make_pptx(workspace / "delivery.pptx")
    removed.unlink()
    candidate, hashes = collect(
        workspace,
        candidate_id="revised",
        final_response="Done",
        initial_hashes=initial,
    )
    assert candidate["environment"]["deleted_files"] == ["obsolete.xlsx"]
    assert candidate["environment"]["file_count"] == 3
    assert [row["path"] for row in candidate["artifact"]["changed_files"]] == [
        "delivery.pptx",
        "report.docx",
    ]
    records = {row["path"]: row for row in workspace_records(evidence_for(candidate))}
    assert records["reference.pdf"]["changed"] is False
    assert records["report.docx"]["changed"] is True
    assert (
        records["report.docx"]["content"]["paragraphs"][-1]
        == "Revised authored conclusion"
    )
    assert hashes["report.docx"] != initial["report.docx"]
    assert hashes["reference.pdf"] == initial["reference.pdf"]
    assert hashes == {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in workspace.iterdir()
    }
    old_records = {
        row["path"]: row for row in workspace_records(evidence_for(original))
    }
    assert (
        "Revised authored conclusion"
        not in old_records["report.docx"]["content"]["paragraphs"]
    )
    assert old_records["report.docx"]["sha256"] == initial["report.docx"]


@pytest.mark.asyncio
async def test_real_docx_content_reassembles_across_source_chunks_and_read_pages(
    tmp_path,
):
    docx = pytest.importorskip("docx", reason="requires the office extra")
    path = tmp_path / "long.docx"
    doc = docx.Document()
    paragraphs = [
        f"Paragraph {i:03d}: " + "authored evidence " * 20 for i in range(160)
    ]
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)
    candidate, _ = collect(
        tmp_path, candidate_id="long", final_response="", initial_hashes={}
    )
    chunks = [
        row for row in candidate["source_documents"] if row["source"] == PRODUCT_LABEL
    ]
    assert len(chunks) > 1 and all(chunk["excerpt_only"] for chunk in chunks)
    assembled = ""
    for chunk in chunks:
        assert chunk["offset"] == len(assembled)
        assert len(chunk["text"]) <= 14000
        assembled += chunk["text"]
    assert len(assembled) == chunks[0]["total_chars"]
    assert json.loads(assembled)[0]["content"]["paragraphs"] == paragraphs
    evidence = evidence_for(candidate)
    (record,) = workspace_records(evidence)
    assert len(record["evidence_refs"]) == len(chunks)
    # Reads after collection must use the frozen evidence, even if the workspace
    # file disappears. The authored PDF/Office parsers have already run for real.
    path.unlink()
    tools = WorkBuddyTools()
    pages, offset, complete = [], 0, False
    for _ in range(10):
        receipt = await tools.run(
            {
                "tool": "workspace.read",
                "arguments": {"paths": ["long.docx"], "offset": offset},
            },
            evidence,
        )
        assert receipt["status"] == "observed"
        result = receipt["result"]
        assert result["offset"] == offset
        assert result["observed_product_refs"] == record["evidence_refs"]
        assert result["excerpt_only"] is True
        pages.append(result["text"])
        if result["next_offset"] is None:
            complete = True
            break
        assert result["next_offset"] == offset + len(result["text"])
        offset = result["next_offset"]
    assert complete and len(pages) > 1
    assert paragraphs[-1] not in pages[0]
    body = json.loads("".join(pages))
    assert body == {"files": [record], "missing_paths": []}
    assert body["files"][0]["content"]["paragraphs"] == paragraphs


def test_pptx_recursively_reads_nested_groups_and_tables_without_flattening(tmp_path):
    pptx = pytest.importorskip("pptx", reason="requires the office extra")
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Inches

    path = tmp_path / "nested.pptx"
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    title.name = "Top-level title"
    title.text = "Original top-level text"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(5), Inches(2))
    table.name = "Nested table"
    rows = [
        ["Metric", "Value"],
        ["Revenue", "1500"],
        ["Status", "Reviewed α\nAccepted"],
    ]
    for row_index, row in enumerate(rows):
        for column_index, text in enumerate(row):
            table.table.cell(row_index, column_index).text = text
    outer = slide.shapes.add_group_shape()
    outer.name = "Outer group"
    inner = outer.shapes.add_group_shape([table])
    inner.name = "Inner group"
    note = inner.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1))
    note.name = "Nested note"
    note.text = "Authored nested conclusion: 证据"
    empty = outer.shapes.add_group_shape()
    empty.name = "Empty group"
    rectangle = outer.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(6), Inches(2), Inches(1), Inches(1)
    )
    rectangle.name = "Unlabeled rectangle"
    deck.save(path)

    result = inspect_file(path)
    shapes = result["slides"][0]["shapes"]
    assert [shape["name"] for shape in shapes] == ["Top-level title", "Outer group"]
    assert shapes[0] == {
        "name": "Top-level title",
        "text": "Original top-level text",
        "left": 914400,
        "top": 914400,
        "width": 4572000,
        "height": 914400,
    }
    descendants = shapes[1]["shapes"]
    assert [shape["name"] for shape in descendants] == [
        "Inner group",
        "Empty group",
        "Unlabeled rectangle",
    ]
    children = descendants[0]["shapes"]
    assert [shape["name"] for shape in children] == ["Nested table", "Nested note"]
    assert children[0]["text"] == "" and children[0]["table"] == rows
    assert children[1]["text"] == "Authored nested conclusion: 证据"
    assert {key: children[1][key] for key in ("left", "top", "width", "height")} == {
        "left": 914400,
        "top": 3657600,
        "width": 4572000,
        "height": 914400,
    }
    assert descendants[1]["text"] == "" and descendants[1]["shapes"] == []
    assert descendants[2]["text"] == ""
    assert "table" not in descendants[2] and "shapes" not in descendants[2]
    assert result["visual_quality"] == "unmeasured"
