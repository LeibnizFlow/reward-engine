"""Inspect actual pre-verifier workspace bytes; never read verifier results."""

import json
import tarfile
import zipfile
from pathlib import Path

from reward_engine.adapters.workbuddy.io import file_hash, write

EXCLUDED = {
    ".git",
    ".opencode",
    "node_modules",
    ".venv",
    "__pycache__",
    ".cache",
    ".pytest_cache",
}
TEXT = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".svg",
    ".py",
    ".js",
    ".ts",
    ".sh",
    ".toml",
    ".ini",
    ".log",
    ".cpp",
    ".h",
    ".cu",
}


def extract_archive(archive, target):
    """Reject all links/special files, including links that escape a snapshot."""
    target = Path(target)
    target.mkdir(parents=True, exist_ok=True)
    with tarfile.open(archive) as bundle:
        for member in bundle:
            path = Path(member.name)
            if (
                path.is_absolute()
                or ".." in path.parts
                or member.issym()
                or member.islnk()
            ):
                raise ValueError("Unsafe workspace archive member")
            if not (member.isfile() or member.isdir()):
                raise ValueError("Special workspace archive member")
            bundle.extract(member, target, filter="data")


def _pptx_shape(shape):
    """Preserve native shape text/geometry and authored table/group contents."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    content = {
        "name": shape.name,
        "text": getattr(shape, "text", ""),
        "left": shape.left,
        "top": shape.top,
        "width": shape.width,
        "height": shape.height,
    }
    if shape.has_table:
        content["table"] = [
            [cell.text for cell in row.cells] for row in shape.table.rows
        ]
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        content["shapes"] = [_pptx_shape(child) for child in shape.shapes]
    return content


def inspect_file(path):
    suffix = path.suffix.lower()
    if suffix in TEXT or suffix == "":
        return {"text": path.read_text(encoding="utf-8", errors="replace")}
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        formulas = load_workbook(path, data_only=False, read_only=True)
        cached = load_workbook(path, data_only=True, read_only=True)
        try:
            return {
                "sheets": [
                    {
                        "name": sheet.title,
                        "cells": [
                            {
                                "cell": cell.coordinate,
                                "value_or_formula": cell.value,
                                "cached_value": cached[sheet.title][
                                    cell.coordinate
                                ].value,
                            }
                            for row in sheet
                            for cell in row
                            if cell.value is not None
                        ],
                    }
                    for sheet in formulas
                ]
            }
        finally:
            formulas.close()
            cached.close()
    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        return {
            "paragraphs": [p.text for p in doc.paragraphs],
            "tables": [
                [[c.text for c in row.cells] for row in t.rows] for t in doc.tables
            ],
        }
    if suffix == ".pptx":
        from pptx import Presentation

        deck = Presentation(path)
        return {
            "slide_count": len(deck.slides),
            "slides": [
                {
                    "slide": i,
                    "shapes": [_pptx_shape(shape) for shape in slide.shapes],
                    "notes": slide.notes_slide.notes_text_frame.text
                    if slide.has_notes_slide
                    else "",
                }
                for i, slide in enumerate(deck.slides, 1)
            ],
            "visual_quality": "unmeasured",
        }
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        return {
            "pages": [
                {"page": i, "text": page.extract_text()}
                for i, page in enumerate(reader.pages, 1)
            ],
            "visual_quality": "unmeasured",
        }
    if suffix in {".zip", ".xmind"}:
        with zipfile.ZipFile(path) as archive:
            return {
                "entries": archive.namelist(),
                "documents": {
                    n: archive.read(n).decode("utf-8", errors="replace")
                    for n in archive.namelist()
                    if Path(n).suffix in {".xml", ".json", ".txt", ".md"}
                    and archive.getinfo(n).file_size < 2_000_000
                },
            }
    return {
        "status": "unknown",
        "reason": "Binary retained; no content or visual interpretation",
    }


def source_chunks(records, label, chunk_chars=14000):
    """Stable, bounded pieces available through core evidence expansion/document.read."""
    text = json.dumps(records, ensure_ascii=False, default=str)
    return [
        {
            "source": label,
            "offset": offset,
            "total_chars": len(text),
            "text": text[offset : offset + chunk_chars],
            "excerpt_only": len(text) > chunk_chars,
        }
        for offset in range(0, len(text), chunk_chars)
    ]


def collect(
    workspace,
    *,
    candidate_id,
    final_response,
    initial_hashes=None,
    output=None,
    public_delivery=False,
):
    workspace = Path(workspace)
    inventory, sources, products = [], [], []
    collected = None
    if public_delivery:
        from reward_engine.adapters.workbuddy.public_delivery import scan_workspace

        collected = scan_workspace(
            workspace, excluded=EXCLUDED, initial_hashes=initial_hashes
        )
    paths = (
        [workspace / r["path"] for r in collected["files"]]
        if collected is not None
        else sorted(workspace.rglob("*"))
    )
    collected_rows = (
        {r["path"]: r for r in collected["files"]} if collected is not None else {}
    )
    for path in paths:
        relative = path.relative_to(workspace)
        if set(relative.parts) & EXCLUDED or not path.is_file() or path.is_symlink():
            if collected is not None:
                raise ValueError("Workspace path changed during collection")
            continue
        name = relative.as_posix()
        row = {"path": name, "sha256": file_hash(path), "bytes": path.stat().st_size}
        row["changed"] = (
            initial_hashes is not None and initial_hashes.get(name) != row["sha256"]
        )
        if collected is not None:
            if row != collected_rows[name]:
                raise ValueError("Workspace bytes changed during collection")
        inventory.append(row)
        try:
            if row["bytes"] > 20_000_000:
                raise ValueError("File exceeds parser budget; original retained")
            content = inspect_file(path)
        except Exception as error:
            content = {
                "status": "unknown",
                "error_type": type(error).__name__,
                "reason": str(error)[:200],
            }
        record = {**row, "content": content}
        (products if row["changed"] else sources).append(record)
    source_documents = source_chunks(sources, "Unmodified public workspace files")
    source_documents += source_chunks(
        products, "Created or modified workspace files (actual delivery evidence)"
    )
    candidate = {
        "candidate_id": candidate_id,
        "artifact": {
            "final_response": final_response,
            "changed_files": [r for r in inventory if r["changed"]],
        },
        "source_documents": source_documents,
        "environment": {
            "workspace": "/workspace",
            "file_count": len(inventory),
            "deleted_files": sorted(
                set(initial_hashes or {}) - {r["path"] for r in inventory}
            ),
            "snapshot_phase": "before_official_verifier",
            "visual_quality": "unmeasured",
        },
    }
    if collected is not None:
        candidate["environment"]["public_delivery_inventory"] = collected
    if output:
        if collected is not None:
            write(Path(output) / "public_delivery_inventory.json", collected)
        write(Path(output) / "inventory.json", inventory)
        write(Path(output) / "candidate.json", candidate)
    return candidate, {r["path"]: r["sha256"] for r in inventory}
